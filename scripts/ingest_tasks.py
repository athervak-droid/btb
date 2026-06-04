"""Ingest the authoring spreadsheet into the benchmark's task + rubric files.

The spreadsheet has two tabs:
  - an "OG set" tab with the full record per task (final_prompt, prompt_context,
    formatting_context, product/workflow_cat/workflow_subcat, and a free-text
    aggregated_rubric_json), and
  - a "shown to agent" tab whose final_prompt is the authoritative, refined
    version the agent actually sees.

It writes, respecting the public/private split:
  - data/tasks.test.jsonl   PUBLIC: task_id + agent-facing final_prompt +
                             product/workflow_cat/workflow_subcat. Nothing else.
  - data/tasks.private.jsonl PRIVATE (gitignored): full records incl.
                             prompt_context / formatting_context + rubric_file.
  - rubrics/<task_id>.json   PRIVATE (gitignored): the parsed rubric, a list of
                             {criterion, category, weight, check?}.

The rubric text is category headers followed by criteria lines each ending in
"(Weight: N)"; we parse those and auto-attach a deterministic check only where a
line clearly matches one (filename, single tab is judge, negatives, gridlines,
formula errors, sourced comments). Everything else routes to the LLM judge.

Usage:
  python scripts/ingest_tasks.py --xlsx "/path/to/data sample for BTB.xlsx"
"""
from __future__ import annotations
import argparse
import json
import os
import re

from openpyxl import load_workbook

KNOWN_CATEGORIES = [
    "Instruction Following", "Technical Correctness", "Transparency & Auditability",
    "Internal Consistency", "Client Readiness",
]
# A single-integer weight marker, in () or [], anywhere in the line:
#   "(Weight: 3)"  "[Weight: 3]"  "(Weight: 3 points)"  "(Weight: 3):"
_WEIGHT = re.compile(r"[\[(]\s*Weight:\s*(\d+)\s*(?:point|pt)?s?\s*[)\]]", re.IGNORECASE)
# A category-header weight, with a range or "each" (NOT a single criterion weight):
_WEIGHT_RANGE = re.compile(r"[\[(]\s*Weight:[^)\]]*(?:-|each)[^)\]]*[)\]]", re.IGNORECASE)


def _strip_weight(line: str) -> str:
    return _WEIGHT_RANGE.sub("", _WEIGHT.sub("", line)).strip(" :.-\t").strip()


def _detect_check(criterion: str) -> dict | None:
    """Attach a deterministic check only for high-confidence, objective lines."""
    c = criterion.lower()
    # output filename (quoted 'x.xlsx' or unquoted "named x.xlsx") -> file_exists
    m = re.search(r"['\"]?([\w\-]+\.(?:xlsx|pptx|docx))['\"]?", criterion)
    if ("saved as" in c or "file naming" in c or "file name" in c
            or "filename" in c or "named" in c or "file execution" in c) and m:
        return {"type": "file_exists", "pattern": m.group(1)}
    if "#ref" in c or "formula error" in c or ("no " in c and "#" in criterion):
        return {"type": "no_formula_errors", "file": "*.xlsx"}
    if "parenthes" in c and "negativ" in c:
        return {"type": "negatives_in_parentheses", "file": "*.xlsx"}
    if "gridline" in c and ("hidden" in c or "off" in c or "removed" in c):
        return {"type": "gridlines_hidden", "file": "*.xlsx"}
    if ("cell comment" in c or "cell note" in c) or ("sourced" in c and "comment" in c):
        return {"type": "has_cell_comments", "file": "*.xlsx", "min_comments": 5}
    return None


def parse_rubric(text: str) -> list[dict]:
    crits: list[dict] = []
    category = "Uncategorized"
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        m = _WEIGHT.search(line)
        if m:  # a criterion: pull weight out, keep the rest as the criterion text
            criterion = _WEIGHT.sub("", line).strip(" :.-\t").strip()
            criterion = re.sub(r"\s+([:;])", r"\1", criterion)  # " :" -> ":" only
            criterion = re.sub(r"\s{2,}", " ", criterion)
            entry = {"criterion": criterion, "category": category, "weight": int(m.group(1))}
            chk = _detect_check(criterion)
            if chk:
                entry["check"] = chk
            crits.append(entry)
            continue
        head = _strip_weight(line)
        if head and (any(head.lower() == k.lower() for k in KNOWN_CATEGORIES)
                     or _WEIGHT_RANGE.search(line)):
            category = head
        elif crits:  # defensive: a wrapped continuation line
            crits[-1]["criterion"] += " " + line
    return crits


def _sheet_by_first_header(wb, want_cols):
    for ws in wb.worksheets:
        hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
        if all(col in hdr for col in want_cols):
            yield ws, hdr


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--og-tab", default="OG set")
    ap.add_argument("--public", default="data/tasks.test.jsonl")
    ap.add_argument("--private", default="data/tasks.private.jsonl")
    ap.add_argument("--rubric-dir", default="rubrics")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    og = wb[args.og_tab]
    oh = [str(c.value).strip() if c.value else "" for c in og[1]]

    # the agent-facing tab is the OTHER one that has task_id + final_prompt
    agent_prompts: dict[str, str] = {}
    for ws, hdr in _sheet_by_first_header(wb, ["task_id", "final_prompt"]):
        if ws.title == args.og_tab:
            continue
        ti, fi = hdr.index("task_id"), hdr.index("final_prompt")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[ti]:
                agent_prompts[str(row[ti]).strip()] = (row[fi] or "").strip()

    def col(name):
        return oh.index(name) if name in oh else None

    ci = {n: col(n) for n in ["task_id", "final_prompt", "prompt_context",
                              "formatting_context", "product", "workflow_cat",
                              "workflow_subcat", "aggregated_rubric_json"]}

    os.makedirs(args.rubric_dir, exist_ok=True)
    os.makedirs(os.path.dirname(args.public) or ".", exist_ok=True)
    public_rows, private_rows = [], []

    for row in og.iter_rows(min_row=2, values_only=True):
        tid = row[ci["task_id"]]
        if not tid:
            continue
        tid = str(tid).strip()
        # authoritative final_prompt comes from the agent tab; fall back to OG
        final_prompt = agent_prompts.get(tid) or (row[ci["final_prompt"]] or "").strip()
        product = (row[ci["product"]] or "").strip() if ci["product"] is not None else ""
        wcat = (row[ci["workflow_cat"]] or "").strip() if ci["workflow_cat"] is not None else ""
        wsub = (row[ci["workflow_subcat"]] or "").strip() if ci["workflow_subcat"] is not None else ""

        rubric = parse_rubric(row[ci["aggregated_rubric_json"]] or "")
        rubric_path = os.path.join(args.rubric_dir, f"{tid}.json")
        with open(rubric_path, "w") as f:
            json.dump(rubric, f, indent=2)

        public_rows.append({"task_id": tid, "final_prompt": final_prompt,
                            "product": product, "workflow_cat": wcat,
                            "workflow_subcat": wsub})
        private_rows.append({"task_id": tid, "final_prompt": final_prompt,
                             "prompt_context": (row[ci["prompt_context"]] or "").strip(),
                             "formatting_context": (row[ci["formatting_context"]] or "").strip(),
                             "product": product, "workflow_cat": wcat,
                             "workflow_subcat": wsub, "rubric_file": rubric_path})

    with open(args.public, "w") as f:
        for r in public_rows:
            f.write(json.dumps(r) + "\n")
    with open(args.private, "w") as f:
        for r in private_rows:
            f.write(json.dumps(r) + "\n")

    print(f"ingested {len(public_rows)} tasks")
    for pub, priv in zip(public_rows, private_rows):
        rub = json.load(open(priv["rubric_file"]))
        checks = sum(1 for c in rub if "check" in c)
        print(f"  {pub['task_id']:6s} {pub['product']:13s} {len(rub):>2} criteria "
              f"({checks} checks / {len(rub)-checks} judge)  prompt={len(pub['final_prompt'])}c")
    print(f"\npublic  -> {args.public}")
    print(f"private -> {args.private} (gitignore this)")
    print(f"rubrics -> {args.rubric_dir}/<task_id>.json (gitignored)")


if __name__ == "__main__":
    main()
