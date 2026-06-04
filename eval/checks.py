"""Deterministic rubric checks against agent deliverables.

Each check returns (score, detail) where score is a float in [0, 1] and detail is
a short human-readable explanation. A rubric criterion routes here when it carries
a `check` block, e.g. {"type": "no_formula_errors", "file": "*.xlsx"}.

These cover the cheap, objective criteria (file present, no errors, formatting).
Subjective / judgment criteria are handled by the LLM judge in score.py instead.
"""
from __future__ import annotations
import glob
import os
import re

try:
    from openpyxl import load_workbook
except ImportError:  # keep import-safe so the module loads without openpyxl
    load_workbook = None

ERROR_STRINGS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!")


def _resolve(outputs_dir: str, pattern: str) -> list[str]:
    return sorted(glob.glob(os.path.join(outputs_dir, "**", pattern), recursive=True))


def _first_xlsx(outputs_dir: str, pattern: str = "*.xlsx") -> str | None:
    hits = _resolve(outputs_dir, pattern)
    return hits[0] if hits else None


def file_exists(outputs_dir, spec):
    pattern = spec.get("pattern", "*")
    hits = _resolve(outputs_dir, pattern)
    if hits:
        return 1.0, f"found {os.path.basename(hits[0])}"
    return 0.0, f"no file matching {pattern}"


def no_formula_errors(outputs_dir, spec):
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path, data_only=True)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(e in cell.value for e in ERROR_STRINGS):
                    bad.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if bad:
        return 0.0, f"{len(bad)} error cell(s): " + ", ".join(bad[:5])
    return 1.0, "no formula errors found"


def negatives_in_parentheses(outputs_dir, spec):
    """Heuristic: of the numeric cells whose number_format encodes negatives,
    what share use a parenthesis format rather than a minus sign."""
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path)
    fmts_with_negative, paren = 0, 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, (int, float)):
                    continue
                fmt = (cell.number_format or "").lower()
                if ";" in fmt or "(" in fmt:  # format defines a negative section
                    fmts_with_negative += 1
                    if "(" in fmt and ")" in fmt:
                        paren += 1
    if fmts_with_negative == 0:
        return 0.5, "no explicit negative formats detected (inconclusive)"
    ratio = paren / fmts_with_negative
    return (1.0 if ratio >= 0.9 else round(ratio, 2)), f"{paren}/{fmts_with_negative} negative formats use parentheses"


def gridlines_hidden(outputs_dir, spec):
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path)
    states = {ws.title: bool(ws.sheet_view.showGridLines) for ws in wb.worksheets}
    hidden = [t for t, shown in states.items() if not shown]
    if hidden:
        return 1.0, f"gridlines hidden on: {', '.join(hidden)}"
    return 0.0, "gridlines visible on all tabs"


def has_cell_comments(outputs_dir, spec):
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path)
    n = sum(1 for ws in wb.worksheets for row in ws.iter_rows()
            for cell in row if cell.comment is not None)
    need = int(spec.get("min_comments", 1))
    return (1.0 if n >= need else round(n / need, 2)), f"{n} cell comment(s) (need >= {need})"


def value_tolerance(outputs_dir, spec):
    """Check a specific sheet!cell against an expected value within a tolerance."""
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path, data_only=True)
    ws = wb[spec["sheet"]] if spec.get("sheet") else wb.worksheets[0]
    val = ws[spec["cell"]].value
    if not isinstance(val, (int, float)):
        return 0.0, f"{spec.get('sheet','')}!{spec['cell']} is not numeric ({val!r})"
    exp, tol = float(spec["expected"]), float(spec.get("tol", 0))
    ok = abs(val - exp) <= tol
    return (1.0 if ok else 0.0), f"{spec['cell']}={val} vs {exp} (+/- {tol})"


def value_tolerance_by_label(outputs_dir, spec):
    """Find a label cell, read the value immediately to its right, check tolerance."""
    path = _first_xlsx(outputs_dir, spec.get("file", "*.xlsx"))
    if not path or load_workbook is None:
        return 0.0, "no workbook found or openpyxl missing"
    wb = load_workbook(path, data_only=True)
    label = str(spec["label"]).strip().lower()
    exp, tol = float(spec["expected"]), float(spec.get("tol", 0))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().lower() == label:
                    right = ws.cell(cell.row, cell.column + 1).value
                    if isinstance(right, (int, float)) and abs(right - exp) <= tol:
                        return 1.0, f"'{label}'={right} vs {exp} (+/- {tol})"
                    return 0.0, f"'{label}' value {right!r} not within tolerance of {exp}"
    return 0.0, f"label '{label}' not found"


DISPATCH = {
    "file_exists": file_exists,
    "no_formula_errors": no_formula_errors,
    "negatives_in_parentheses": negatives_in_parentheses,
    "gridlines_hidden": gridlines_hidden,
    "has_cell_comments": has_cell_comments,
    "value_tolerance": value_tolerance,
    "value_tolerance_by_label": value_tolerance_by_label,
}


def run_check(spec: dict, outputs_dir: str):
    """Route a check spec to its function. Returns (score_float, detail_str)."""
    fn = DISPATCH.get(spec.get("type"))
    if fn is None:
        return 0.0, f"unknown check type: {spec.get('type')!r}"
    try:
        return fn(outputs_dir, spec)
    except Exception as e:  # a check failing shouldn't crash the whole run
        return 0.0, f"check error: {e}"
